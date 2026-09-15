"""Workbook contents — the deliverable most users keep."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from app.models import Lead
from app.services.export import build_csv, build_workbook, export_filename


def make_lead(**overrides: object) -> Lead:
    defaults: dict[str, object] = {
        "id": uuid.uuid4(),
        "job_id": uuid.uuid4(),
        "task_id": uuid.uuid4(),
        "image_id": uuid.uuid4(),
        "first_name": "Priya",
        "last_name": "Raghavan",
        "position": "Head of Business Development",
        "company": "Meridian Logistics",
        "location": "Mumbai, India",
        "phone": "+919876543210",
        "email": "priya.raghavan@meridianlog.com",
        "created_at": datetime.now(UTC),
    }
    return Lead(**{**defaults, **overrides})


def sheet(leads: list[Lead]):
    job_id = uuid.uuid4()
    data = build_workbook(leads, job_id=job_id)
    return load_workbook(io.BytesIO(data))


class TestWorkbook:
    def test_phone_is_stored_as_text(self) -> None:
        """Excel strips the leading plus from a numeric cell.

        Without a text format, +919876543210 opens as 919876543210 and the
        number is no longer dialable internationally.
        """
        ws = sheet([make_lead()])["Leads"]
        assert ws.cell(row=2, column=7).number_format == "@"
        assert ws.cell(row=2, column=7).value == "+919876543210"

    def test_email_is_a_mailto_link(self) -> None:
        ws = sheet([make_lead()])["Leads"]
        cell = ws.cell(row=2, column=8)
        assert cell.hyperlink is not None
        assert cell.hyperlink.target == "mailto:priya.raghavan@meridianlog.com"

    def test_header_is_frozen_and_filterable(self) -> None:
        ws = sheet([make_lead()])["Leads"]
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None

    def test_a_doubtful_field_is_shaded(self) -> None:
        ws = sheet([make_lead(confidence={"company": 0.6})])["Leads"]
        assert ws.cell(row=2, column=5).fill.fgColor.rgb == "00FFF2CC"

    def test_a_confident_field_is_not_shaded(self) -> None:
        ws = sheet([make_lead(confidence={"company": 1.0})])["Leads"]
        assert ws.cell(row=2, column=5).fill.fgColor.rgb != "00FFF2CC"

    def test_an_absent_field_is_empty_not_invented(self) -> None:
        ws = sheet([make_lead(position=None)])["Leads"]
        assert ws.cell(row=2, column=4).value is None


class TestSummary:
    def test_a_single_doubtful_field_counts_the_row(self) -> None:
        """Counted per field, not by row mean.

        Six confident fields and one doubtful one average above the threshold,
        so a mean-based count claimed "0 flagged" while the sheet visibly
        shaded a cell — contradicting itself.
        """
        leads = [
            make_lead(
                confidence=dict.fromkeys(
                    ("first_name", "last_name", "company", "location", "phone", "email"), 1.0
                )
                | {"position": 0.0}
            )
        ]
        rows = dict(sheet(leads)["Summary"].iter_rows(values_only=True))
        assert rows["Rows with a field to check"] == 1

    def test_a_fully_confident_row_is_not_counted(self) -> None:
        leads = [
            make_lead(
                confidence=dict.fromkeys(
                    (
                        "first_name",
                        "last_name",
                        "position",
                        "company",
                        "location",
                        "phone",
                        "email",
                    ),
                    1.0,
                )
            )
        ]
        rows = dict(sheet(leads)["Summary"].iter_rows(values_only=True))
        assert rows["Rows with a field to check"] == 0


class TestCsv:
    def test_a_bom_is_written_for_excel(self) -> None:
        """Without it Excel renders non-ASCII names as mojibake."""
        assert build_csv([make_lead()]).startswith(b"\xef\xbb\xbf")

    def test_non_ascii_names_survive_a_round_trip(self) -> None:
        data = build_csv([make_lead(first_name="Hiroshi", last_name="田中")])
        assert "田中" in data.decode("utf-8-sig")


@pytest.mark.parametrize("extension", ["xlsx", "csv"])
def test_filename_carries_the_batch_and_date(extension: str) -> None:
    job_id = uuid.uuid4()
    name = export_filename(job_id, extension)
    assert name.startswith(f"leads_{str(job_id)[:8]}_")
    assert name.endswith(f".{extension}")
